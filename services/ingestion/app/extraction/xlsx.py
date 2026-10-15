from __future__ import annotations

from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import date, datetime, time
from decimal import Decimal, InvalidOperation
import hashlib
from pathlib import Path
import re
from time import perf_counter
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter, range_boundaries
from openpyxl.worksheet.worksheet import Worksheet

from ..vision.base import SlideReviewProvider, VisionProvider, VisionProviderError
from ..vision.prompt import XLSX_PROMPT_VERSION, XLSX_REVIEW_PROMPT_VERSION
from .common import DocumentBuilder
from .errors import ExtractionFailure
from .models import (
    ExtractionQuality,
    NormalizedDocument,
    NormalizedWorksheet,
    SlideBoundingBox,
    SlideReviewDecision,
    SlideReviewDecisionFinding,
    SlideReviewProvenance,
    SlideVisionProvenance,
    SourceLocator,
    WorkbookCell,
    WorkbookDefinedName,
    WorkbookMetadata,
    WorksheetNativeElement,
    WorksheetRegion,
)
from .pptx_reconcile import reconcile_worksheet
from .pptx_review import decide_worksheet_review, failed_worksheet_review
from .xlsx_renderer import RenderedWorksheet, XlsxRenderer
from .xlsx_validation import (
    XlsxPackageInventory,
    XlsxPackageLimits,
    validate_xlsx_package,
)


_FORMULA_ERROR = re.compile(r"^#(?:REF!|VALUE!|NAME\?|DIV/0!|N/A|NUM!|NULL!)$")
_EXTERNAL_FORMULA = re.compile(r"\[[^\]]+\]")


def extract_xlsx(
    path: Path,
    source_filename: str,
    content_hash: str,
    *,
    content_profile: str = "general_document",
    renderer: XlsxRenderer | None = None,
    vision_provider: VisionProvider | None = None,
    review_provider: SlideReviewProvider | None = None,
    vision_prompt_version: str = XLSX_PROMPT_VERSION,
    review_prompt_version: str = XLSX_REVIEW_PROMPT_VERSION,
    max_concurrency: int = 4,
    package_limits: XlsxPackageLimits | None = None,
    visual_mode: str = "selective",
) -> NormalizedDocument:
    limits = package_limits or XlsxPackageLimits()
    inventory = validate_xlsx_package(path, limits)
    builder = DocumentBuilder(source_filename, "xlsx", content_hash)
    try:
        formula_book = load_workbook(
            path,
            read_only=False,
            data_only=False,
            keep_links=True,
            rich_text=True,
        )
        cached_book = load_workbook(
            path,
            read_only=False,
            data_only=True,
            keep_links=False,
            rich_text=True,
        )
    except Exception as exc:
        raise ExtractionFailure(
            "CORRUPT_DOCUMENT", "XLSX file could not be opened"
        ) from exc

    try:
        defined_names = _defined_names(formula_book)
        named_ranges = _named_ranges_by_sheet(formula_book)
        worksheets: list[NormalizedWorksheet] = []
        total_cells = 0
        empty_sheets: list[str] = []
        for worksheet_index, worksheet in enumerate(
            formula_book.worksheets, start=1
        ):
            cached_sheet = cached_book[worksheet.title]
            cells = _cells(worksheet, cached_sheet)
            total_cells += len(cells)
            if total_cells > limits.max_cells:
                raise ExtractionFailure(
                    "XLSX_LIMIT_EXCEEDED",
                    f"XLSX contains more than {limits.max_cells} non-empty cells",
                )
            if not cells:
                empty_sheets.append(worksheet.title)
                continue
            worksheets.append(
                _native_worksheet(
                    worksheet_index,
                    worksheet,
                    cells,
                    named_ranges.get(worksheet.title, []),
                    builder,
                    inventory,
                    visual_mode,
                )
            )
        if empty_sheets:
            builder.warnings.append(
                "worksheets without values: " + ", ".join(empty_sheets)
            )
        native_document = builder.build().model_copy(
            update={
                "worksheets": worksheets,
                "workbook_metadata": _workbook_metadata(
                    formula_book, inventory, defined_names
                ),
                "content_profile": content_profile,
            }
        )
    finally:
        formula_book.close()
        cached_book.close()

    selected = [
        sheet.worksheet_index for sheet in worksheets if sheet.selection_reasons
    ]
    if not selected:
        return native_document.model_copy(
            update={"worksheets": [_auto_pass(sheet) for sheet in worksheets]}
        )
    if renderer is None:
        raise ExtractionFailure(
            "XLSX_RENDERER_REQUIRED",
            "XLSX visual evidence requires the LibreOffice worksheet renderer",
            partial_document=native_document,
        )
    if vision_provider is None or review_provider is None:
        raise ExtractionFailure(
            "XLSX_VISION_PROVIDER_REQUIRED",
            "XLSX visual evidence requires extraction and review providers",
            partial_document=native_document,
        )

    try:
        rendered_result = renderer.render(path, selected)
    except ExtractionFailure as exc:
        raise ExtractionFailure(
            exc.code, str(exc), partial_document=native_document
        ) from exc
    rendered_by_index = {
        item.worksheet_index: item for item in rendered_result.worksheets
    }
    missing = sorted(set(selected) - set(rendered_by_index))
    if missing:
        raise ExtractionFailure(
            "XLSX_RENDER_FAILED",
            "renderer omitted worksheets: " + ", ".join(map(str, missing)),
            partial_document=native_document,
        )

    recalculated = _recalculated_values(
        rendered_result.recalculated_workbook, worksheets
    )
    prepared = [
        _with_recalculation(
            sheet,
            recalculated.get(sheet.title, {}),
            rendered_by_index.get(sheet.worksheet_index),
            bool(inventory.external_links),
            bool(inventory.embedded_objects),
        )
        for sheet in worksheets
    ]
    selected_sheets = [sheet for sheet in prepared if sheet.selection_reasons]
    completed: dict[int, NormalizedWorksheet] = {}
    failures: list[str] = []
    workers = max(1, min(max_concurrency, len(selected_sheets)))
    with ThreadPoolExecutor(max_workers=workers) as executor:
        futures = {
            executor.submit(
                _enrich_worksheet,
                sheet,
                rendered_by_index[sheet.worksheet_index],
                content_profile,
                vision_provider,
                review_provider,
                vision_prompt_version,
                review_prompt_version,
                renderer.version,
            ): sheet.worksheet_index
            for sheet in selected_sheets
        }
        for future in as_completed(futures):
            worksheet_index = futures[future]
            try:
                completed[worksheet_index] = future.result()
            except (VisionProviderError, OSError, ValueError) as exc:
                failures.append(
                    f"worksheet {worksheet_index}: {type(exc).__name__}: {exc}"
                )

    normalized_sheets = [
        completed.get(sheet.worksheet_index, _auto_pass(sheet))
        if sheet.selection_reasons
        else _auto_pass(sheet)
        for sheet in prepared
    ]
    warnings = list(native_document.quality.warnings)
    warnings.extend(
        f"worksheet {sheet.worksheet_index} render was truncated to the configured page limit"
        for sheet in normalized_sheets
        if "rendered_pages_truncated" in sheet.selection_reasons
    )
    normalized = native_document.model_copy(
        update={
            "worksheets": normalized_sheets,
            "quality": ExtractionQuality(
                status="WARNING" if warnings else "PASS",
                block_count=len(native_document.blocks),
                text_character_count=native_document.quality.text_character_count,
                warnings=list(dict.fromkeys(warnings)),
            ),
        }
    )
    if failures:
        raise ExtractionFailure(
            "XLSX_VISION_FAILED",
            "XLSX worksheet processing failed for " + "; ".join(sorted(failures)),
            partial_document=normalized,
        )
    material = list(
        dict.fromkeys(
            f"worksheet {sheet.worksheet_index} [{finding.severity}]: "
            f"{finding.description}"
            for sheet in normalized.worksheets
            if sheet.review is not None
            and sheet.review.status == "REVIEW_REQUIRED"
            for finding in sheet.review.findings
        )
    )
    if material:
        raise ExtractionFailure(
            "XLSX_REVIEW_REQUIRED",
            "material worksheet discrepancies require review: "
            + "; ".join(material),
            partial_document=normalized,
        )
    return normalized


def _cells(worksheet: Worksheet, cached_sheet: Worksheet) -> list[WorkbookCell]:
    merged_by_coordinate = {
        cell.coordinate: str(merged)
        for merged in worksheet.merged_cells.ranges
        for row in worksheet[str(merged)]
        for cell in row
    }
    cells = []
    for row in worksheet.iter_rows():
        for cell in row:
            if cell.value is None:
                continue
            formula = str(cell.value) if cell.data_type == "f" else None
            cached = cached_sheet[cell.coordinate].value if formula else None
            cells.append(
                WorkbookCell(
                    coordinate=cell.coordinate,
                    row=cell.row,
                    column=cell.column,
                    value=_json_value(cell.value),
                    data_type=str(cell.data_type),
                    number_format=cell.number_format or None,
                    formula=formula,
                    cached_value=_json_value(cached),
                    formula_status=(
                        "NOT_RECALCULATED" if formula else "NOT_FORMULA"
                    ),
                    merged_range=merged_by_coordinate.get(cell.coordinate),
                )
            )
    return cells


def _native_worksheet(
    worksheet_index: int,
    worksheet: Worksheet,
    cells: list[WorkbookCell],
    named_ranges: list[tuple[str, str]],
    builder: DocumentBuilder,
    inventory: XlsxPackageInventory,
    visual_mode: str,
) -> NormalizedWorksheet:
    builder.heading(
        worksheet.title,
        SourceLocator(worksheet=worksheet.title),
        level=1,
    )
    regions = _regions(worksheet, cells, named_ranges)
    populated_regions = []
    for region in regions:
        start = len(builder.blocks)
        min_col, min_row, max_col, max_row = range_boundaries(region.cell_range)
        rows = _cells_by_rows(
            cells, min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col
        )
        for row_number, row_cells in rows:
            display = [_display_value(cell) for cell in row_cells]
            first_col = min(cell.column for cell in row_cells)
            last_col = max(cell.column for cell in row_cells)
            builder.add(
                (
                    "table_header"
                    if region.header_row is not None
                    and row_number == region.header_row
                    else "table_row"
                ),
                "\t".join(display),
                SourceLocator(
                    worksheet=worksheet.title,
                    cell_range=(
                        f"{get_column_letter(first_col)}{row_number}:"
                        f"{get_column_letter(last_col)}{row_number}"
                    ),
                ),
            )
        block_ids = [block.id for block in builder.blocks[start:]]
        populated_regions.append(region.model_copy(update={"block_ids": block_ids}))

    native_markdown = _worksheet_markdown(worksheet.title, populated_regions)
    native_elements = _native_elements(worksheet, populated_regions, cells)
    reasons = _selection_reasons(
        worksheet,
        cells,
        populated_regions,
        inventory,
        visual_mode,
    )
    return NormalizedWorksheet(
        worksheet_index=worksheet_index,
        title=worksheet.title,
        state=worksheet.sheet_state,
        used_range=worksheet.calculate_dimension(),
        freeze_panes=(
            str(worksheet.freeze_panes) if worksheet.freeze_panes is not None else None
        ),
        auto_filter_range=worksheet.auto_filter.ref or None,
        merged_ranges=[str(item) for item in worksheet.merged_cells.ranges],
        table_names=list(worksheet.tables.keys()),
        chart_count=len(worksheet._charts),
        image_count=len(worksheet._images),
        pivot_table_count=len(getattr(worksheet, "_pivots", [])),
        conditional_formatting_count=len(worksheet.conditional_formatting),
        data_validation_count=len(worksheet.data_validations.dataValidation),
        cells=cells,
        regions=populated_regions,
        native_block_ids=[
            block_id for region in populated_regions for block_id in region.block_ids
        ],
        native_markdown=native_markdown,
        native_elements=native_elements,
        selection_reasons=reasons,
        combined_markdown=native_markdown,
    )


def _regions(
    worksheet: Worksheet,
    cells: list[WorkbookCell],
    named_ranges: list[tuple[str, str]],
) -> list[WorksheetRegion]:
    occupied = {(cell.row, cell.column): cell for cell in cells}
    claimed: set[tuple[int, int]] = set()
    regions: list[WorksheetRegion] = []

    def add_region(kind: str, name: str | None, cell_range: str) -> None:
        try:
            min_col, min_row, max_col, max_row = range_boundaries(cell_range)
        except ValueError:
            return
        present = [
            cell
            for (row, column), cell in occupied.items()
            if min_row <= row <= max_row and min_col <= column <= max_col
        ]
        if not present or all(
            (cell.row, cell.column) in claimed for cell in present
        ):
            return
        rows = sorted({cell.row for cell in present})
        header_row = _header_row(present, rows[0])
        columns = [
            _display_value(occupied[(header_row, column)])
            for column in range(min_col, max_col + 1)
            if header_row is not None and (header_row, column) in occupied
        ]
        actual_min_col = min(cell.column for cell in present)
        actual_max_col = max(cell.column for cell in present)
        actual_min_row = min(cell.row for cell in present)
        actual_max_row = max(cell.row for cell in present)
        canonical_range = (
            f"{get_column_letter(actual_min_col)}{actual_min_row}:"
            f"{get_column_letter(actual_max_col)}{actual_max_row}"
        )
        region_rows = _cells_by_rows(
            present,
            min_row=actual_min_row,
            max_row=actual_max_row,
            min_col=actual_min_col,
            max_col=actual_max_col,
        )
        regions.append(
            WorksheetRegion(
                id=f"ws-{worksheet.title}-region-{len(regions) + 1:04d}",
                kind=kind,
                name=name,
                cell_range=canonical_range,
                header_row=header_row,
                columns=columns,
                row_count=len(region_rows),
                column_count=actual_max_col - actual_min_col + 1,
                native_markdown=_region_markdown(region_rows, header_row),
            )
        )
        claimed.update((cell.row, cell.column) for cell in present)

    for table in worksheet.tables.values():
        add_region("excel_table", table.displayName, table.ref)
    for name, cell_range in named_ranges:
        add_region("named_range", name, cell_range)

    remaining_rows = sorted(
        {
            cell.row
            for cell in cells
            if (cell.row, cell.column) not in claimed
        }
    )
    for row_group in _consecutive_groups(remaining_rows):
        group_cells = [
            cell
            for cell in cells
            if cell.row in row_group and (cell.row, cell.column) not in claimed
        ]
        columns = sorted({cell.column for cell in group_cells})
        for column_group in _consecutive_groups(columns):
            component = [cell for cell in group_cells if cell.column in column_group]
            if not component:
                continue
            cell_range = (
                f"{get_column_letter(min(cell.column for cell in component))}"
                f"{min(cell.row for cell in component)}:"
                f"{get_column_letter(max(cell.column for cell in component))}"
                f"{max(cell.row for cell in component)}"
            )
            add_region("detected_table", None, cell_range)
    return regions


def _cells_by_rows(
    cells: list[WorkbookCell],
    *,
    min_row: int,
    max_row: int,
    min_col: int,
    max_col: int,
) -> list[tuple[int, list[WorkbookCell]]]:
    lookup = {(cell.row, cell.column): cell for cell in cells}
    rows = []
    for row in range(min_row, max_row + 1):
        row_cells = [
            lookup[(row, column)]
            for column in range(min_col, max_col + 1)
            if (row, column) in lookup
        ]
        if row_cells:
            rows.append((row, row_cells))
    return rows


def _header_row(cells: list[WorkbookCell], first_row: int) -> int | None:
    first = [cell for cell in cells if cell.row == first_row]
    if not first:
        return None
    values = [_display_value(cell) for cell in first]
    if all(cell.formula is None and isinstance(cell.value, str) for cell in first):
        if len(set(value.casefold() for value in values)) == len(values):
            return first_row
    return None


def _consecutive_groups(values: list[int]) -> list[list[int]]:
    if not values:
        return []
    groups = [[values[0]]]
    for value in values[1:]:
        if value == groups[-1][-1] + 1:
            groups[-1].append(value)
        else:
            groups.append([value])
    return groups


def _region_markdown(
    rows: list[tuple[int, list[WorkbookCell]]], header_row: int | None
) -> str:
    if not rows:
        return ""
    width = max(len(row_cells) for _, row_cells in rows)
    rendered = []
    for _, row_cells in rows:
        values = [_display_value(cell).replace("|", "\\|") for cell in row_cells]
        values.extend([""] * (width - len(values)))
        rendered.append("| " + " | ".join(values) + " |")
    if header_row is not None:
        rendered.insert(1, "| " + " | ".join(["---"] * width) + " |")
    return "\n".join(rendered)


def _worksheet_markdown(title: str, regions: list[WorksheetRegion]) -> str:
    lines = [f"# Worksheet: {title}"]
    for region in regions:
        label = region.name or region.kind.replace("_", " ").title()
        lines.extend(
            [
                "",
                f"## {label} ({region.cell_range})",
                "",
                region.native_markdown,
            ]
        )
    return "\n".join(lines).strip()


def _native_elements(
    worksheet: Worksheet,
    regions: list[WorksheetRegion],
    cells: list[WorkbookCell],
) -> list[WorksheetNativeElement]:
    max_row = max(cell.row for cell in cells)
    max_col = max(cell.column for cell in cells)
    elements = []
    for order, region in enumerate(regions):
        min_col, min_row, range_max_col, range_max_row = range_boundaries(
            region.cell_range
        )
        elements.append(
            WorksheetNativeElement(
                element_id=f"worksheet-region-{order + 1}",
                kind="table",
                content=region.native_markdown,
                bbox=_cell_bbox(
                    min_col,
                    min_row,
                    range_max_col,
                    range_max_row,
                    max_col,
                    max_row,
                ),
                block_ids=region.block_ids,
                reading_order=order,
                attributes={
                    "region_id": region.id,
                    "region_kind": region.kind,
                    "cell_range": region.cell_range,
                    "name": region.name,
                },
            )
        )
    for chart_index, chart in enumerate(worksheet._charts, start=1):
        elements.append(
            WorksheetNativeElement(
                element_id=f"worksheet-chart-{chart_index}",
                kind="chart",
                content=_chart_content(chart),
                reading_order=len(elements),
                attributes={"chart_type": type(chart).__name__},
            )
        )
    for image_index, _image in enumerate(worksheet._images, start=1):
        elements.append(
            WorksheetNativeElement(
                element_id=f"worksheet-image-{image_index}",
                kind="image",
                content=f"Embedded worksheet image {image_index}",
                reading_order=len(elements),
            )
        )
    return elements


def _cell_bbox(
    min_col: int,
    min_row: int,
    max_col: int,
    max_row: int,
    sheet_max_col: int,
    sheet_max_row: int,
) -> SlideBoundingBox:
    x = round((min_col - 1) / max(sheet_max_col, 1) * 1000)
    y = round((min_row - 1) / max(sheet_max_row, 1) * 1000)
    right = round(max_col / max(sheet_max_col, 1) * 1000)
    bottom = round(max_row / max(sheet_max_row, 1) * 1000)
    return SlideBoundingBox(
        x=min(1000, x),
        y=min(1000, y),
        width=max(0, min(1000 - x, right - x)),
        height=max(0, min(1000 - y, bottom - y)),
    )


def _chart_content(chart: Any) -> str:
    references = []
    for series in getattr(chart, "ser", []):
        for attribute in ("val", "cat", "xVal", "yVal"):
            source = getattr(series, attribute, None)
            for reference_name in ("numRef", "strRef"):
                reference = getattr(source, reference_name, None)
                formula = getattr(reference, "f", None)
                if formula:
                    references.append(str(formula))
    suffix = (
        "; source ranges: " + ", ".join(dict.fromkeys(references))
        if references
        else ""
    )
    return f"{type(chart).__name__}{suffix}"


def _selection_reasons(
    worksheet: Worksheet,
    cells: list[WorkbookCell],
    regions: list[WorksheetRegion],
    inventory: XlsxPackageInventory,
    visual_mode: str,
) -> list[str]:
    if visual_mode == "all":
        return ["visual_mode_all"]
    flags = (
        (worksheet._charts, "charts"),
        (worksheet._images, "images"),
        (worksheet.merged_cells.ranges, "merged_cells"),
        (worksheet.conditional_formatting, "conditional_formatting"),
        (getattr(worksheet, "_pivots", []), "pivot_tables"),
        ([cell for cell in cells if cell.formula], "formulas"),
        ([region for region in regions if region.header_row is None], "ambiguous_headers"),
        (worksheet.sheet_state != "visible", "hidden_worksheet"),
        (inventory.external_links, "external_links"),
        (inventory.embedded_objects, "embedded_objects"),
    )
    return [reason for present, reason in flags if present]


def _recalculated_values(
    recalculated_path: Path,
    worksheets: list[NormalizedWorksheet],
) -> dict[str, dict[str, Any]]:
    try:
        workbook = load_workbook(
            recalculated_path,
            read_only=True,
            data_only=True,
            keep_links=False,
        )
    except Exception as exc:
        raise ExtractionFailure(
            "XLSX_RECALCULATION_FAILED",
            "LibreOffice recalculated workbook could not be opened",
        ) from exc
    try:
        result = {}
        for worksheet in worksheets:
            if worksheet.title not in workbook.sheetnames:
                continue
            sheet = workbook[worksheet.title]
            result[worksheet.title] = {
                cell.coordinate: _json_value(sheet[cell.coordinate].value)
                for cell in worksheet.cells
                if cell.formula is not None
            }
        return result
    finally:
        workbook.close()


def _with_recalculation(
    worksheet: NormalizedWorksheet,
    recalculated: dict[str, Any],
    rendered: RenderedWorksheet | None,
    has_external_links: bool,
    has_embedded_objects: bool,
) -> NormalizedWorksheet:
    updated_cells = []
    reasons = list(worksheet.selection_reasons)
    for cell in worksheet.cells:
        if cell.formula is None:
            updated_cells.append(cell)
            continue
        recalculated_value = recalculated.get(cell.coordinate)
        if _is_formula_error(cell.cached_value) or _is_formula_error(recalculated_value):
            status = "ERROR"
        elif recalculated_value is None:
            status = "NOT_RECALCULATED"
        elif cell.cached_value is None:
            status = "MISSING_CACHE"
        elif _equivalent(cell.cached_value, recalculated_value):
            status = "MATCH"
        else:
            status = "VALUE_CHANGED"
        updated_cells.append(
            cell.model_copy(
                update={
                    "recalculated_value": recalculated_value,
                    "formula_status": status,
                }
            )
        )
    if rendered is not None and rendered.pages_truncated:
        reasons.append("rendered_pages_truncated")
    if has_external_links and "external_links" not in reasons:
        reasons.append("external_links")
    if has_embedded_objects and "embedded_objects" not in reasons:
        reasons.append("embedded_objects")
    return worksheet.model_copy(
        update={
            "cells": updated_cells,
            "selection_reasons": list(dict.fromkeys(reasons)),
        }
    )


def _deterministic_findings(
    worksheet: NormalizedWorksheet,
    rendered: RenderedWorksheet,
) -> list[SlideReviewDecisionFinding]:
    findings = []
    for cell in worksheet.cells:
        if cell.formula_status == "VALUE_CHANGED":
            findings.append(
                _finding(
                    "ATOMIC_FACT_CONFLICT",
                    "CRITICAL",
                    f"Formula result changed at {cell.coordinate}: original cached "
                    f"value {cell.cached_value!r}, isolated recalculation "
                    f"{cell.recalculated_value!r}.",
                    "The original XLSX cache and LibreOffice recalculation disagree; "
                    "the original source was not modified.",
                )
            )
        elif cell.formula_status == "ERROR":
            findings.append(
                _finding(
                    "ATOMIC_FACT_CONFLICT",
                    "CRITICAL",
                    f"Formula error detected at {cell.coordinate}.",
                    f"Formula {cell.formula!r}; cached={cell.cached_value!r}; "
                    f"recalculated={cell.recalculated_value!r}.",
                )
            )
        elif cell.formula_status == "MISSING_CACHE":
            findings.append(
                _finding(
                    "PARTIAL_OR_UNCERTAIN",
                    "MEDIUM",
                    f"Formula {cell.coordinate} has no original cached value.",
                    f"LibreOffice recalculated {cell.recalculated_value!r}, but this "
                    "supporting result cannot silently replace absent native evidence.",
                )
            )
        elif cell.formula_status == "NOT_RECALCULATED":
            findings.append(
                _finding(
                    "PARTIAL_OR_UNCERTAIN",
                    "MEDIUM",
                    f"Formula {cell.coordinate} could not be independently recalculated.",
                    f"Formula expression: {cell.formula!r}.",
                )
            )
        if cell.formula and _EXTERNAL_FORMULA.search(cell.formula):
            findings.append(
                _finding(
                    "MISSING_KEY_ELEMENT",
                    "HIGH",
                    f"Formula {cell.coordinate} depends on an external workbook.",
                    f"External formula: {cell.formula!r}; external sources were not opened.",
                )
            )
    if "ambiguous_headers" in worksheet.selection_reasons:
        findings.append(
            _finding(
                "PARTIAL_OR_UNCERTAIN",
                "MEDIUM",
                "At least one detected data region has no reliable header row.",
                "The region was detected from occupied cells rather than an Excel "
                "table or named range.",
            )
        )
    if "embedded_objects" in worksheet.selection_reasons:
        findings.append(
            _finding(
                "MISSING_KEY_ELEMENT",
                "HIGH",
                "The workbook contains an embedded object that was not executed.",
                "Embedded OOXML package parts are quarantined as untrusted evidence.",
            )
        )
    if rendered.pages_truncated:
        findings.append(
            _finding(
                "PARTIAL_OR_UNCERTAIN",
                "MEDIUM",
                "The worksheet visual preview exceeded the configured page limit.",
                f"The worksheet rendered to {rendered.page_count} pages; the contact "
                "sheet contains only the configured maximum.",
            )
        )
    return findings


def _finding(
    category: str, severity: str, description: str, evidence: str
) -> SlideReviewDecisionFinding:
    return SlideReviewDecisionFinding(
        category=category,
        severity=severity,
        source="deterministic",
        description=description,
        evidence=evidence,
    )


def _enrich_worksheet(
    worksheet: NormalizedWorksheet,
    rendered: RenderedWorksheet,
    content_profile: str,
    vision_provider: VisionProvider,
    review_provider: SlideReviewProvider,
    vision_prompt_version: str,
    review_prompt_version: str,
    render_version: str,
) -> NormalizedWorksheet:
    started = perf_counter()
    vision = vision_provider.analyze_worksheet(
        rendered.path,
        worksheet_index=worksheet.worksheet_index,
        worksheet_title=worksheet.title,
        content_profile=content_profile,
        prompt_version=vision_prompt_version,
    )
    vision_duration_ms = max(0, round((perf_counter() - started) * 1000))
    vision_provenance = SlideVisionProvenance(
        provider=vision_provider.name,
        model=vision_provider.model,
        prompt_version=vision_prompt_version,
        response_hash=hashlib.sha256(
            vision.model_dump_json().encode("utf-8")
        ).hexdigest(),
        duration_ms=vision_duration_ms,
    )
    reconciliation = reconcile_worksheet(
        native_markdown=worksheet.native_markdown,
        native_elements=worksheet.native_elements,
        vision=vision,
    )
    review_started = perf_counter()
    try:
        judge_result = review_provider.review_worksheet(
            rendered.path,
            worksheet_index=worksheet.worksheet_index,
            worksheet_title=worksheet.title,
            content_profile=content_profile,
            native_markdown=worksheet.native_markdown,
            native_elements=[
                element.model_dump(mode="json")
                for element in worksheet.native_elements
            ],
            vision=vision,
            reconciliation=reconciliation,
            prompt_version=review_prompt_version,
        )
        judge_provenance = SlideReviewProvenance(
            provider=review_provider.name,
            model=review_provider.model,
            prompt_version=review_prompt_version,
            response_hash=hashlib.sha256(
                judge_result.model_dump_json().encode("utf-8")
            ).hexdigest(),
            duration_ms=max(0, round((perf_counter() - review_started) * 1000)),
        )
        review = decide_worksheet_review(
            vision=vision,
            reconciliation=reconciliation,
            judge_result=judge_result,
            judge_provenance=judge_provenance,
            deterministic_findings=_deterministic_findings(worksheet, rendered),
        )
    except VisionProviderError as exc:
        review = failed_worksheet_review(
            f"Independent worksheet judge failed: {exc}"
        )
    return worksheet.model_copy(
        update={
            "image_path": rendered.relative_path,
            "image_hash": rendered.content_hash,
            "rendered_page_count": rendered.page_count,
            "render_version": render_version,
            "vision": vision,
            "vision_provenance": vision_provenance,
            "reconciliation": reconciliation,
            "review": review,
            "combined_markdown": reconciliation.retrieval_markdown,
        }
    )


def _auto_pass(worksheet: NormalizedWorksheet) -> NormalizedWorksheet:
    if worksheet.review is not None:
        return worksheet
    return worksheet.model_copy(
        update={
            "review": SlideReviewDecision(
                version="xlsx-worksheet-review-v1",
                status="AUTO_PASS",
                severity="NONE",
                findings=[],
            )
        }
    )


def _defined_names(workbook: Any) -> list[WorkbookDefinedName]:
    result = []
    for name in workbook.defined_names.values():
        destinations = []
        try:
            destinations = [
                f"{sheet}!{coordinate}" for sheet, coordinate in name.destinations
            ]
        except (AttributeError, TypeError, ValueError):
            pass
        result.append(
            WorkbookDefinedName(
                name=name.name,
                destinations=destinations,
                expression=name.attr_text if not destinations else None,
                scope=str(name.localSheetId) if name.localSheetId is not None else None,
            )
        )
    return result


def _named_ranges_by_sheet(workbook: Any) -> dict[str, list[tuple[str, str]]]:
    result: dict[str, list[tuple[str, str]]] = {}
    for name in workbook.defined_names.values():
        try:
            destinations = list(name.destinations)
        except (AttributeError, TypeError, ValueError):
            continue
        for sheet, coordinate in destinations:
            if sheet in workbook.sheetnames and ":" in coordinate:
                result.setdefault(sheet, []).append((name.name, coordinate))
    return result


def _workbook_metadata(
    workbook: Any,
    inventory: XlsxPackageInventory,
    defined_names: list[WorkbookDefinedName],
) -> WorkbookMetadata:
    external_links = [*inventory.external_links]
    for link in getattr(workbook, "_external_links", []):
        target = getattr(getattr(link, "file_link", None), "Target", None)
        if target:
            external_links.append(str(target))
    return WorkbookMetadata(
        date_system="1904" if "1904" in str(workbook.epoch) else "1900",
        calculation_mode=getattr(workbook.calculation, "calcMode", None),
        defined_names=defined_names,
        external_links=list(dict.fromkeys(external_links)),
        embedded_objects=list(inventory.embedded_objects),
        package_entry_count=inventory.entry_count,
        package_uncompressed_bytes=inventory.uncompressed_bytes,
        package_compression_ratio=inventory.compression_ratio,
    )


def _display_value(cell: WorkbookCell) -> str:
    value = (
        cell.cached_value
        if cell.formula is not None and cell.cached_value is not None
        else cell.value
    )
    return "" if value is None else str(value).strip()


def _json_value(value: Any) -> Any:
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    if isinstance(value, Decimal):
        return str(value)
    if isinstance(value, (str, int, float, bool)) or value is None:
        return value
    return str(value)


def _is_formula_error(value: Any) -> bool:
    return isinstance(value, str) and bool(_FORMULA_ERROR.fullmatch(value.strip()))


def _equivalent(left: Any, right: Any) -> bool:
    if left == right:
        return True
    try:
        left_number = Decimal(str(left))
        right_number = Decimal(str(right))
        tolerance = max(Decimal("1e-9"), abs(left_number) * Decimal("1e-9"))
        return abs(left_number - right_number) <= tolerance
    except (InvalidOperation, ValueError):
        return str(left).strip() == str(right).strip()
