from __future__ import annotations

import hashlib
import json
import time
from concurrent.futures import ThreadPoolExecutor, as_completed
from dataclasses import dataclass
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from ..vision.base import SlideReviewProvider, VisionProvider, VisionProviderError
from ..vision.prompt import PROMPT_VERSION, REVIEW_PROMPT_VERSION
from .common import DocumentBuilder, clean_text
from .errors import ExtractionFailure
from .models import (
    ExtractionQuality,
    NormalizedDocument,
    NormalizedSlide,
    SlideBoundingBox,
    SlideNativeElement,
    SlideVisionProvenance,
    SlideReviewProvenance,
    SlideVisionResult,
    SourceLocator,
)
from .pptx_link import link_slides
from .pptx_reconcile import reconcile_slide
from .pptx_review import decide_slide_review, failed_slide_review
from .pptx_renderer import PptxRenderer, RenderedSlide


@dataclass(frozen=True)
class NativeSlide:
    slide_number: int
    title: str | None
    block_ids: list[str]
    native_markdown: str
    native_elements: list[SlideNativeElement]
    speaker_notes: str | None


def extract_native_slide_elements(
    path: Path, source_filename: str, content_hash: str
) -> dict[int, list[SlideNativeElement]]:
    """Rebuild native PPTX element geometry without rendering or model calls."""

    try:
        presentation = Presentation(path)
    except Exception as exc:
        raise ExtractionFailure(
            "CORRUPT_DOCUMENT", "PPTX file could not be opened"
        ) from exc
    builder = DocumentBuilder(source_filename, "pptx", content_hash)
    slides = _extract_native_slides(presentation, builder)
    return {slide.slide_number: slide.native_elements for slide in slides}


def extract_pptx(
    path: Path,
    source_filename: str,
    content_hash: str,
    *,
    content_profile: str,
    renderer: PptxRenderer | None,
    vision_provider: VisionProvider | None,
    review_provider: SlideReviewProvider | None,
    vision_prompt_version: str = PROMPT_VERSION,
    review_prompt_version: str = REVIEW_PROMPT_VERSION,
    max_concurrency: int = 4,
) -> NormalizedDocument:
    try:
        presentation = Presentation(path)
    except Exception as exc:
        raise ExtractionFailure(
            "CORRUPT_DOCUMENT", "PPTX file could not be opened"
        ) from exc

    builder = DocumentBuilder(source_filename, "pptx", content_hash)
    native_slides = _extract_native_slides(presentation, builder)
    native_document = _build_native_document(builder)

    if renderer is None:
        raise ExtractionFailure(
            "PPTX_RENDERER_REQUIRED",
            "PPTX hybrid extraction requires a configured slide renderer",
            partial_document=native_document,
        )
    if vision_provider is None:
        raise ExtractionFailure(
            "PPTX_VISION_PROVIDER_REQUIRED",
            "PPTX hybrid extraction requires a configured vision provider",
            partial_document=native_document,
        )
    if review_provider is None:
        raise ExtractionFailure(
            "PPTX_REVIEW_PROVIDER_REQUIRED",
            "PPTX hybrid extraction requires a configured independent review provider",
            partial_document=native_document,
        )

    try:
        rendered = renderer.render(path, len(native_slides))
    except ExtractionFailure as exc:
        raise ExtractionFailure(
            exc.code,
            str(exc),
            partial_document=native_document,
        ) from exc

    rendered_by_number = {item.slide_number: item for item in rendered}
    completed: dict[int, NormalizedSlide] = {}
    failures: list[str] = []
    workers = max(1, min(max_concurrency, len(native_slides)))
    with ThreadPoolExecutor(max_workers=workers) as executor:
        futures = {
            executor.submit(
                _enrich_slide,
                native,
                rendered_by_number[native.slide_number],
                content_profile,
                vision_provider,
                review_provider,
                vision_prompt_version,
                review_prompt_version,
                renderer.version,
            ): native.slide_number
            for native in native_slides
        }
        for future in as_completed(futures):
            slide_number = futures[future]
            try:
                completed[slide_number] = future.result()
            except (VisionProviderError, OSError, ValueError) as exc:
                failures.append(
                    f"slide {slide_number}: {type(exc).__name__}: {exc}"
                )

    slides = [completed[number] for number in sorted(completed)]
    warnings = list(native_document.quality.warnings)
    for slide in slides:
        warnings.extend(
            f"slide {slide.slide_number} uncertainty: {item}"
            for item in slide.vision.uncertainties
        )
    document = _with_slides(native_document, slides, warnings)

    if failures:
        raise ExtractionFailure(
            "PPTX_VISION_FAILED",
            "visual extraction failed for " + "; ".join(sorted(failures)),
            partial_document=document,
        )

    material = list(dict.fromkeys(
        f"slide {slide.slide_number} [{finding.severity}]: {finding.description}"
        for slide in slides
        if slide.review is not None and slide.review.status == "REVIEW_REQUIRED"
        for finding in slide.review.findings
    ))
    if material:
        raise ExtractionFailure(
            "PPTX_REVIEW_REQUIRED",
            "material native/visual discrepancies require review: "
            + "; ".join(material),
            partial_document=document,
        )
    return document


def _extract_native_slides(
    presentation: Presentation, builder: DocumentBuilder
) -> list[NativeSlide]:
    native_slides: list[NativeSlide] = []
    empty_slides: list[int] = []
    table_number = 0
    for slide_number, slide in enumerate(presentation.slides, start=1):
        builder.section_path = []
        before = len(builder.blocks)
        native_elements: list[SlideNativeElement] = []
        title_shape = slide.shapes.title
        title_shape_id = title_shape.shape_id if title_shape is not None else None
        title = None
        if title_shape is not None and title_shape.text.strip():
            title = clean_text(title_shape.text)
            element_before = len(builder.blocks)
            builder.heading(
                title_shape.text,
                SourceLocator(slide=slide_number),
                level=1,
            )
            native_elements.append(
                SlideNativeElement(
                    element_id=_shape_element_id(
                        slide_number, title_shape.shape_id
                    ),
                    kind="title",
                    content=f"# {title}",
                    bbox=_shape_bbox(
                        title_shape,
                        presentation.slide_width,
                        presentation.slide_height,
                    ),
                    block_ids=[
                        block.id for block in builder.blocks[element_before:]
                    ],
                )
            )
        for shape in slide.shapes:
            if title_shape_id is not None and shape.shape_id == title_shape_id:
                continue
            if getattr(shape, "has_table", False):
                table_number += 1
                markdown = _table_markdown(shape.table)
                element_before = len(builder.blocks)
                builder.add(
                    "table",
                    markdown,
                    SourceLocator(slide=slide_number, table=table_number),
                )
                native_elements.append(
                    SlideNativeElement(
                        element_id=_shape_element_id(
                            slide_number, shape.shape_id
                        ),
                        kind="table",
                        content=markdown,
                        bbox=_shape_bbox(
                            shape,
                            presentation.slide_width,
                            presentation.slide_height,
                        ),
                        block_ids=[
                            block.id for block in builder.blocks[element_before:]
                        ],
                    )
                )
            elif getattr(shape, "has_text_frame", False):
                element_before = len(builder.blocks)
                paragraphs: list[str] = []
                for paragraph in shape.text_frame.paragraphs:
                    if paragraph.text.strip():
                        paragraphs.append(clean_text(paragraph.text))
                        builder.add(
                            "paragraph",
                            paragraph.text,
                            SourceLocator(slide=slide_number),
                        )
                if paragraphs:
                    native_elements.append(
                        SlideNativeElement(
                            element_id=_shape_element_id(
                                slide_number, shape.shape_id
                            ),
                            kind="text",
                            content="\n\n".join(paragraphs),
                            bbox=_shape_bbox(
                                shape,
                                presentation.slide_width,
                                presentation.slide_height,
                            ),
                            block_ids=[
                                block.id
                                for block in builder.blocks[element_before:]
                            ],
                        )
                    )
            elif getattr(shape, "has_chart", False):
                chart_title = ""
                try:
                    if shape.chart.has_title:
                        chart_title = clean_text(
                            shape.chart.chart_title.text_frame.text
                        )
                except (AttributeError, ValueError):
                    chart_title = ""
                native_elements.append(
                    SlideNativeElement(
                        element_id=_shape_element_id(
                            slide_number, shape.shape_id
                        ),
                        kind="chart",
                        content=chart_title,
                        bbox=_shape_bbox(
                            shape,
                            presentation.slide_width,
                            presentation.slide_height,
                        ),
                    )
                )
            elif shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                native_elements.append(
                    SlideNativeElement(
                        element_id=_shape_element_id(
                            slide_number, shape.shape_id
                        ),
                        kind="image",
                        content="",
                        bbox=_shape_bbox(
                            shape,
                            presentation.slide_width,
                            presentation.slide_height,
                        ),
                    )
                )
        try:
            notes = clean_text(slide.notes_slide.notes_text_frame.text)
        except (AttributeError, ValueError):
            notes = ""
        visible_blocks = list(builder.blocks[before:])
        if notes:
            element_before = len(builder.blocks)
            builder.add(
                "slide_note", notes, SourceLocator(slide=slide_number)
            )
            native_elements.append(
                SlideNativeElement(
                    element_id=f"slide-{slide_number:04d}-notes",
                    kind="notes",
                    content=notes,
                    block_ids=[
                        block.id for block in builder.blocks[element_before:]
                    ],
                )
            )
        slide_blocks = list(builder.blocks[before:])
        if not slide_blocks:
            empty_slides.append(slide_number)
        native_slides.append(
            NativeSlide(
                slide_number=slide_number,
                title=title,
                block_ids=[block.id for block in slide_blocks],
                native_markdown=_blocks_markdown(visible_blocks),
                native_elements=native_elements,
                speaker_notes=notes or None,
            )
        )
    if empty_slides:
        builder.warnings.append(
            "slides without native extractable text: "
            + ", ".join(str(slide) for slide in empty_slides)
        )
    return native_slides


def _shape_element_id(slide_number: int, shape_id: int) -> str:
    return f"slide-{slide_number:04d}-shape-{shape_id:04d}"


def _shape_bbox(shape, slide_width: int, slide_height: int) -> SlideBoundingBox:
    def scale(value: int, total: int) -> int:
        return min(1000, max(0, round(value / total * 1000)))

    return SlideBoundingBox(
        x=scale(shape.left, slide_width),
        y=scale(shape.top, slide_height),
        width=scale(shape.width, slide_width),
        height=scale(shape.height, slide_height),
    )


def _table_markdown(table) -> str:
    rows = [
        [_escape_markdown_cell(cell.text) for cell in row.cells]
        for row in table.rows
    ]
    if not rows:
        return ""
    width = max(len(row) for row in rows)
    rows = [row + [""] * (width - len(row)) for row in rows]
    header = rows[0]
    lines = [
        "| " + " | ".join(header) + " |",
        "| " + " | ".join("---" for _ in header) + " |",
    ]
    lines.extend("| " + " | ".join(row) + " |" for row in rows[1:])
    return "\n".join(lines)


def _escape_markdown_cell(value: str) -> str:
    return clean_text(value).replace("|", "\\|").replace("\n", "<br>")


def _blocks_markdown(blocks) -> str:
    parts: list[str] = []
    for block in blocks:
        if block.kind == "heading":
            parts.append(f"# {block.text}")
        else:
            parts.append(block.text)
    return "\n\n".join(parts).strip()


def _build_native_document(builder: DocumentBuilder) -> NormalizedDocument:
    if builder.blocks:
        return builder.build()
    return NormalizedDocument(
        source_filename=builder.source_filename,
        source_format=builder.source_format,
        content_hash=builder.content_hash,
        blocks=[],
        quality=ExtractionQuality(
            status="WARNING",
            block_count=0,
            text_character_count=0,
            warnings=["presentation has no native extractable text"],
        ),
    )


def _enrich_slide(
    native: NativeSlide,
    rendered: RenderedSlide,
    content_profile: str,
    provider: VisionProvider,
    review_provider: SlideReviewProvider,
    prompt_version: str,
    review_prompt_version: str,
    render_version: str,
) -> NormalizedSlide:
    started = time.perf_counter()
    result = provider.analyze_slide(
        rendered.path,
        slide_number=native.slide_number,
        content_profile=content_profile,
        native_markdown=native.native_markdown,
        native_elements=[
            element.model_dump(mode="json") for element in native.native_elements
        ],
        prompt_version=prompt_version,
    )
    duration_ms = round((time.perf_counter() - started) * 1000)
    response_payload = result.model_dump(mode="json")
    response_hash = hashlib.sha256(
        json.dumps(
            response_payload, sort_keys=True, separators=(",", ":")
        ).encode("utf-8")
    ).hexdigest()
    reconciliation = reconcile_slide(
        native_markdown=native.native_markdown,
        native_elements=native.native_elements,
        speaker_notes=native.speaker_notes,
        vision=result,
    )
    native_element_payload = [
        element.model_dump(mode="json") for element in native.native_elements
    ]
    try:
        review_started = time.perf_counter()
        judge_result = review_provider.review_slide(
            rendered.path,
            slide_number=native.slide_number,
            content_profile=content_profile,
            native_markdown=native.native_markdown,
            native_elements=native_element_payload,
            vision=result,
            reconciliation=reconciliation,
            prompt_version=review_prompt_version,
        )
        review_duration_ms = round((time.perf_counter() - review_started) * 1000)
        review_payload = judge_result.model_dump(mode="json")
        review_hash = hashlib.sha256(
            json.dumps(
                review_payload, sort_keys=True, separators=(",", ":")
            ).encode("utf-8")
        ).hexdigest()
        review = decide_slide_review(
            vision=result,
            reconciliation=reconciliation,
            judge_result=judge_result,
            judge_provenance=SlideReviewProvenance(
                provider=review_provider.name,
                model=review_provider.model,
                prompt_version=review_prompt_version,
                response_hash=review_hash,
                duration_ms=review_duration_ms,
            ),
        )
    except (VisionProviderError, OSError, ValueError) as exc:
        review = failed_slide_review(
            f"Independent slide review failed: {type(exc).__name__}: {exc}"
        )
    return NormalizedSlide(
        slide_number=native.slide_number,
        title=native.title,
        native_block_ids=native.block_ids,
        native_markdown=native.native_markdown,
        native_elements=native.native_elements,
        speaker_notes=native.speaker_notes,
        image_path=rendered.relative_path,
        image_hash=rendered.content_hash,
        render_version=render_version,
        vision=result,
        vision_provenance=SlideVisionProvenance(
            provider=provider.name,
            model=provider.model,
            prompt_version=prompt_version,
            response_hash=response_hash,
            duration_ms=duration_ms,
        ),
        reconciliation=reconciliation,
        review=review,
        combined_markdown=reconciliation.retrieval_markdown,
    )


def _with_slides(
    native: NormalizedDocument,
    slides: list[NormalizedSlide],
    warnings: list[str],
) -> NormalizedDocument:
    quality = native.quality.model_copy(
        update={
            "status": "WARNING" if warnings else "PASS",
            "warnings": list(dict.fromkeys(warnings)),
        }
    )
    return native.model_copy(
        update={"slides": link_slides(slides), "quality": quality}
    )
