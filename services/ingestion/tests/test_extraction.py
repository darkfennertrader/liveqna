from __future__ import annotations

import hashlib
import os
import shutil
from pathlib import Path

import pytest
from docx import Document
from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.table import Table, TableStyleInfo
from PIL import Image
from pptx import Presentation
from pptx.util import Inches
from pypdf import PdfWriter
from reportlab.pdfgen import canvas

from app.extraction.errors import ExtractionFailure
from app.extraction.models import (
    SlideBoundingBox,
    SlideDiscrepancy,
    SlideVisionResult,
    SlideVisualElement,
    SlideReviewFinding,
    SlideReviewResult,
)
from app.extraction.pipeline import ExtractionPipeline
from app.extraction.pptx_report import (
    build_pptx_llm_report,
    build_pptx_review_report,
)
from app.extraction.pdf_renderer import (
    PyMuPdfRenderer,
    RenderedPage,
)
from app.extraction.pdf_report import build_pdf_llm_report
from app.extraction.xlsx_report import build_xlsx_review_report
from app.extraction.pptx_renderer import (
    LibreOfficePptxRenderer,
    RenderedSlide,
)
from app.extraction.xlsx_renderer import (
    RenderedWorksheet,
    XlsxRenderResult,
)


PIPELINE = ExtractionPipeline()


class FakeRenderer:
    version = "fake-renderer-v1"

    def render(self, source: Path, expected_slides: int) -> list[RenderedSlide]:
        directory = source.parent / "slides"
        directory.mkdir(exist_ok=True)
        rendered = []
        for slide_number in range(1, expected_slides + 1):
            path = directory / f"slide-{slide_number:04d}.png"
            path.write_bytes(f"fake-png-{slide_number}".encode())
            rendered.append(
                RenderedSlide(
                    slide_number=slide_number,
                    path=path,
                    relative_path=f"slides/{path.name}",
                    content_hash=digest(path),
                )
            )
        return rendered


class FakePdfRenderer:
    version = "fake-pdf-renderer-v1"

    def render(
        self, source: Path, expected_pages: int, page_numbers: list[int]
    ) -> list[RenderedPage]:
        assert expected_pages >= max(page_numbers)
        directory = source.parent / "pages"
        directory.mkdir(exist_ok=True)
        rendered = []
        for page_number in page_numbers:
            path = directory / f"page-{page_number:04d}.png"
            path.write_bytes(f"fake-png-{page_number}".encode())
            rendered.append(
                RenderedPage(
                    page_number=page_number,
                    path=path,
                    relative_path=f"pages/{path.name}",
                    content_hash=digest(path),
                )
            )
        return rendered


class FakeXlsxRenderer:
    version = "fake-xlsx-renderer-v1"

    def render(
        self, source: Path, worksheet_indexes: list[int]
    ) -> XlsxRenderResult:
        directory = source.parent / "worksheets"
        directory.mkdir(exist_ok=True)
        recalculated = directory / "recalculated.xlsx"
        workbook = load_workbook(source, data_only=False)
        for worksheet in workbook.worksheets:
            for row in worksheet.iter_rows():
                for cell in row:
                    if cell.data_type == "f":
                        cell.value = 3
        workbook.save(recalculated)
        workbook.close()
        rendered = []
        source_workbook = load_workbook(source, read_only=True)
        try:
            for worksheet_index in worksheet_indexes:
                title = source_workbook.worksheets[worksheet_index - 1].title
                path = directory / f"worksheet-{worksheet_index:04d}.png"
                path.write_bytes(f"fake-xlsx-{worksheet_index}".encode())
                rendered.append(
                    RenderedWorksheet(
                        worksheet_index=worksheet_index,
                        title=title,
                        path=path,
                        relative_path=f"worksheets/{path.name}",
                        content_hash=digest(path),
                        page_count=1,
                        pages_truncated=False,
                    )
                )
        finally:
            source_workbook.close()
        return XlsxRenderResult(
            worksheets=rendered,
            recalculated_workbook=recalculated,
        )


class FakeVisionProvider:
    name = "fake"
    model = "fake-vision-v1"

    def __init__(self, *, material_discrepancy: bool = False) -> None:
        self.material_discrepancy = material_discrepancy

    def analyze_slide(self, image_path: Path, **kwargs) -> SlideVisionResult:
        assert image_path.is_file()
        return self._result()

    def analyze_page(self, image_path: Path, **kwargs) -> SlideVisionResult:
        assert image_path.is_file()
        return self._result()

    def analyze_worksheet(self, image_path: Path, **kwargs) -> SlideVisionResult:
        assert image_path.is_file()
        return self._result()

    def review_slide(self, image_path: Path, **kwargs) -> SlideReviewResult:
        assert image_path.is_file()
        return SlideReviewResult(
            summary="Independent comparison completed.",
            findings=[
                SlideReviewFinding(
                    category="NO_DISCREPANCY",
                    description="No discrepancy found.",
                    evidence="Rendered, native, and reconciled evidence agree.",
                )
            ],
        )

    def review_page(self, image_path: Path, **kwargs) -> SlideReviewResult:
        return self.review_slide(image_path, **kwargs)

    def review_worksheet(self, image_path: Path, **kwargs) -> SlideReviewResult:
        return self.review_slide(image_path, **kwargs)

    def _result(self) -> SlideVisionResult:
        return SlideVisionResult(
            transcription_markdown="Agent architecture\n\nOne orchestrator",
            tables_markdown=[
                "| Layer | Purpose |\n| --- | --- |\n| Registry | Lifecycle |"
            ],
            visual_description="A governed orchestration architecture.",
            elements=[
                SlideVisualElement(
                    kind="text",
                    content="Agent architecture\n\nOne orchestrator",
                    bbox=SlideBoundingBox(x=40, y=40, width=920, height=300),
                ),
                SlideVisualElement(
                    kind="diagram",
                    content="A governed orchestration architecture.",
                    bbox=SlideBoundingBox(x=100, y=350, width=800, height=500),
                ),
            ],
            labels=["agent architecture", "registry"],
            discrepancies=(
                [
                    SlideDiscrepancy(
                        description="A material wording mismatch",
                        material=True,
                    )
                ]
                if self.material_discrepancy
                else []
            ),
            uncertainties=[],
        )


def hybrid_pipeline(*, material_discrepancy: bool = False) -> ExtractionPipeline:
    return ExtractionPipeline(
        pdf_renderer=FakePdfRenderer(),
        pptx_renderer=FakeRenderer(),
        xlsx_renderer=FakeXlsxRenderer(),
        vision_provider=FakeVisionProvider(
            material_discrepancy=material_discrepancy
        ),
        review_provider=FakeVisionProvider(
            material_discrepancy=material_discrepancy
        ),
        vision_max_concurrency=2,
    )


def digest(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def extract(path: Path, document_type: str = "official"):
    return PIPELINE.extract(path, path.name, document_type, digest(path))


def test_markdown_extraction_is_deterministic_and_preserves_structure(
    tmp_path: Path,
) -> None:
    path = tmp_path / "architecture.md"
    path.write_text(
        "# Target architecture\n\nClient-facing overview.\n\n"
        "## Components\n\n| Service | Role |\n|---|---|\n| API | Entry |\n",
        encoding="utf-8",
    )

    first = extract(path)
    second = extract(path)

    assert first == second
    assert [block.kind for block in first.blocks] == [
        "heading",
        "paragraph",
        "heading",
        "table",
    ]
    assert first.blocks[-1].source.line_start == 7
    assert first.blocks[-1].section_path == [
        "Target architecture",
        "Components",
    ]


def test_text_transcript_retains_speakers_timestamps_and_lines(
    tmp_path: Path,
) -> None:
    path = tmp_path / "meeting.txt"
    path.write_text(
        "[00:01:05] Alice: We approved the target design.\n"
        "00:01:12 Bob: I will record the decision.\n",
        encoding="utf-8",
    )

    normalized = extract(path, document_type="transcript")

    assert [block.kind for block in normalized.blocks] == [
        "transcript_turn",
        "transcript_turn",
    ]
    assert normalized.blocks[0].source.timestamp == "00:01:05"
    assert normalized.blocks[0].source.speaker == "Alice"
    assert normalized.blocks[1].source.line_start == 2


def test_docx_extraction_preserves_headings_paragraphs_and_tables(
    tmp_path: Path,
) -> None:
    path = tmp_path / "decision.docx"
    document = Document()
    document.add_heading("Architecture decision", level=1)
    document.add_paragraph("Use an isolated client project.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Component"
    table.cell(0, 1).text = "Owner"
    table.cell(1, 0).text = "Registry"
    table.cell(1, 1).text = "Architecture"
    document.save(path)

    normalized = extract(path)

    assert [block.kind for block in normalized.blocks] == [
        "heading",
        "paragraph",
        "table",
    ]
    assert normalized.blocks[0].source.paragraph == 1
    assert normalized.blocks[-1].source.table == 1
    assert "Component\tOwner" in normalized.blocks[-1].text


def test_xlsx_extraction_retains_worksheet_and_cell_ranges(tmp_path: Path) -> None:
    path = tmp_path / "services.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Services"
    sheet.append(["Service", "Cloud"])
    sheet.append(["Qdrant", "Local"])
    workbook.create_sheet("Empty")
    workbook.save(path)

    normalized = extract(path)

    assert [block.kind for block in normalized.blocks] == [
        "heading",
        "table_header",
        "table_row",
    ]
    assert normalized.blocks[1].source.worksheet == "Services"
    assert normalized.blocks[1].source.cell_range == "A1:B1"
    assert normalized.quality.status == "WARNING"
    assert normalized.quality.warnings == ["worksheets without values: Empty"]
    assert normalized.schema_version == "1.7"
    assert len(normalized.worksheets) == 1
    assert normalized.worksheets[0].review.status == "AUTO_PASS"
    assert normalized.worksheets[0].regions[0].cell_range == "A1:B2"


def test_xlsx_declared_tables_take_precedence_over_detected_regions(
    tmp_path: Path,
) -> None:
    path = tmp_path / "tables.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Services"
    sheet.append(["Service", "Cloud"])
    sheet.append(["Qdrant", "Local"])
    table = Table(displayName="ServicesTable", ref="A1:B2")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2")
    sheet.add_table(table)
    workbook.save(path)

    normalized = extract(path)

    worksheet = normalized.worksheets[0]
    assert worksheet.table_names == ["ServicesTable"]
    assert len(worksheet.regions) == 1
    assert worksheet.regions[0].kind == "excel_table"
    assert worksheet.regions[0].name == "ServicesTable"


def test_xlsx_formula_without_native_cache_requires_human_review(
    tmp_path: Path,
) -> None:
    path = tmp_path / "formula.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Totals"
    sheet.append(["Amount", "Total"])
    sheet.append([1, "=SUM(A2:A3)"])
    sheet.append([2, None])
    workbook.save(path)

    with pytest.raises(ExtractionFailure) as raised:
        hybrid_pipeline().extract(
            path, path.name, "official", digest(path), "financial_model"
        )

    assert raised.value.code == "XLSX_REVIEW_REQUIRED"
    assert raised.value.partial_document is not None
    worksheet = raised.value.partial_document.worksheets[0]
    formula = next(cell for cell in worksheet.cells if cell.formula)
    assert formula.cached_value is None
    assert formula.recalculated_value == 3
    assert formula.formula_status == "MISSING_CACHE"
    assert worksheet.image_path == "worksheets/worksheet-0001.png"
    assert worksheet.review.status == "REVIEW_REQUIRED"
    report = build_xlsx_review_report(raised.value.partial_document)
    assert "## Worksheet 1: Totals" in report
    assert "Formula audit" in report
    assert "Human resolution questions" in report


def test_pptx_extraction_retains_slide_text_tables_and_notes(
    tmp_path: Path,
) -> None:
    path = tmp_path / "architecture.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "Agent architecture"
    slide.placeholders[1].text = "One orchestrator with governed tools"
    table_shape = slide.shapes.add_table(
        2, 2, Inches(1), Inches(3), Inches(6), Inches(1.5)
    )
    table_shape.table.cell(0, 0).text = "Layer"
    table_shape.table.cell(0, 1).text = "Purpose"
    table_shape.table.cell(1, 0).text = "Registry"
    table_shape.table.cell(1, 1).text = "Lifecycle"
    slide.notes_slide.notes_text_frame.text = "Validate with the client."
    presentation.save(path)

    normalized = hybrid_pipeline().extract(
        path, path.name, "artifact", digest(path), "general_presentation"
    )

    kinds = [block.kind for block in normalized.blocks]
    assert kinds == ["heading", "paragraph", "table", "slide_note"]
    assert all(block.source.slide == 1 for block in normalized.blocks)
    assert normalized.blocks[2].source.table == 1
    assert normalized.blocks[2].text.startswith("| Layer | Purpose |")
    assert len(normalized.slides) == 1
    normalized_slide = normalized.slides[0]
    assert normalized.schema_version == "1.7"
    assert normalized_slide.image_path == "slides/slide-0001.png"
    assert normalized_slide.vision_provenance.provider == "fake"
    assert normalized_slide.native_elements
    assert normalized_slide.native_elements[0].bbox is not None
    assert normalized_slide.reconciliation is not None
    assert normalized_slide.review is not None
    assert normalized_slide.review.status == "AUTO_PASS"
    assert normalized_slide.reconciliation.suppressed_duplicate_visual_ids == [
        "visual-element-1"
    ]
    assert normalized_slide.reconciliation.included_visual_ids == [
        "visual-element-2"
    ]
    assert "## Unmatched visual text" not in normalized_slide.combined_markdown
    assert "## Visual evidence" in normalized_slide.combined_markdown
    assert "## Speaker notes" in normalized_slide.combined_markdown


def test_pptx_llm_report_exposes_model_output_and_reconciliation(
    tmp_path: Path,
) -> None:
    path = tmp_path / "architecture.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "Agent architecture"
    slide.placeholders[1].text = "One orchestrator with governed tools"
    presentation.save(path)

    normalized = hybrid_pipeline().extract(
        path, path.name, "artifact", digest(path), "general_presentation"
    )
    report = build_pptx_llm_report(normalized)

    assert "## What to compare" in report
    assert "## Slide 1: Agent architecture" in report
    assert "![Rendered slide 1](slides/slide-0001.png)" in report
    assert "### LLM visual description" in report
    assert "A governed orchestration architecture." in report
    assert "### Evidence reconciliation" in report
    assert "Suppressed duplicate visual evidence" in report
    assert "### Human validation for this slide" in report
    assert "### Final reconciled output" in report
    assert "### Provenance" in report
    assert "- Provider: `fake`" in report
    assert "### Independent review decision" in report
    assert "- Status: **AUTO_PASS**" in report
    assert "## Automated review summary" in report
    assert "`NO_DISCREPANCY`" in report
    assert "- Category: `NO_DISCREPANCY`" in report
    assert "- Reason: No discrepancy found." in report

    exception_report = build_pptx_review_report(normalized)
    assert "Slides requiring review: **0**" in exception_report
    assert "No slide was placed in the human-review queue." in exception_report


def test_pptx_material_discrepancy_preserves_partial_output(
    tmp_path: Path,
) -> None:
    path = tmp_path / "review.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "Review required"
    slide.placeholders[1].text = "Native wording"
    presentation.save(path)

    with pytest.raises(ExtractionFailure) as raised:
        hybrid_pipeline(material_discrepancy=True).extract(
            path, path.name, "artifact", digest(path), "general_presentation"
        )

    assert raised.value.code == "PPTX_REVIEW_REQUIRED"
    assert raised.value.partial_document is not None
    assert len(raised.value.partial_document.slides) == 1
    reconciliation = raised.value.partial_document.slides[0].reconciliation
    assert reconciliation is not None
    assert reconciliation.conflicts == ["A material wording mismatch"]


def test_pptx_requires_visual_pipeline_and_preserves_native_output(
    tmp_path: Path,
) -> None:
    path = tmp_path / "native-only.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "Native title"
    presentation.save(path)

    with pytest.raises(ExtractionFailure) as raised:
        extract(path)

    assert raised.value.code == "PPTX_RENDERER_REQUIRED"
    assert raised.value.partial_document is not None
    assert raised.value.partial_document.blocks[0].text == "Native title"


@pytest.mark.skipif(
    os.environ.get("RUN_LIBREOFFICE_TESTS") != "1"
    or shutil.which("soffice") is None
    or shutil.which("pdftoppm") is None,
    reason="set RUN_LIBREOFFICE_TESTS=1 with LibreOffice and pdftoppm available",
)
def test_libreoffice_renderer_creates_one_stable_png_per_slide(
    tmp_path: Path,
) -> None:
    path = tmp_path / "render.pptx"
    presentation = Presentation()
    for number in range(1, 3):
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        slide.shapes.title.text = f"Slide {number}"
    presentation.save(path)

    renderer = LibreOfficePptxRenderer(timeout_seconds=60)
    first = renderer.render(path, 2)
    second = renderer.render(path, 2)

    assert [item.slide_number for item in first] == [1, 2]
    assert all(item.path.is_file() for item in first)
    assert [item.content_hash for item in first] == [
        item.content_hash for item in second
    ]


def test_pdf_extraction_retains_page_numbers(tmp_path: Path) -> None:
    path = tmp_path / "architecture.pdf"
    document = canvas.Canvas(str(path))
    document.drawString(72, 750, "Architecture overview")
    document.showPage()
    document.drawString(72, 750, "Second page decision")
    document.save()

    normalized = hybrid_pipeline().extract(
        path, path.name, "official", digest(path), "general_document"
    )

    assert {block.source.page for block in normalized.blocks} == {1, 2}
    assert [page.page_number for page in normalized.pages] == [1, 2]
    assert all(page.native_elements for page in normalized.pages)
    assert all(page.reconciliation is not None for page in normalized.pages)
    assert all(page.review.status == "AUTO_PASS" for page in normalized.pages)
    assert normalized.quality.status == "PASS"


def test_pdf_transcript_retains_source_metadata_without_inventing_turns(
    tmp_path: Path,
) -> None:
    path = tmp_path / "transcript.pdf"
    document = canvas.Canvas(str(path))
    document.drawString(72, 760, "Transcription")
    document.drawString(72, 740, "Video 1: meeting.mp4")
    document.drawString(
        72,
        720,
        "Time Range: 00:00:00 to 00:59:48 Duration: 3588.88 seconds",
    )
    document.drawString(72, 700, "Continuous prose with no speaker labels.")
    document.save()

    normalized = extract(path, document_type="transcript")

    assert normalized.transcript_metadata is not None
    assert normalized.transcript_metadata.video_name == "meeting.mp4"
    assert normalized.transcript_metadata.time_range_start == "00:00:00"
    assert normalized.transcript_metadata.time_range_end == "00:59:48"
    assert normalized.transcript_metadata.duration_seconds == 3588.88
    assert all(block.kind == "paragraph" for block in normalized.blocks)
    assert all(block.source.speaker is None for block in normalized.blocks)
    assert all(block.source.timestamp is None for block in normalized.blocks)


def test_every_pdf_page_is_independently_enriched_and_reviewed(tmp_path: Path) -> None:
    path = tmp_path / "architecture.pdf"
    image_path = tmp_path / "diagram.png"
    Image.new("RGB", (1200, 700), color="white").save(image_path)
    document = canvas.Canvas(str(path))
    document.drawString(72, 760, "Architecture diagram")
    document.drawImage(str(image_path), 72, 200, width=468, height=273)
    document.save()

    normalized = hybrid_pipeline().extract(
        path, path.name, "official", digest(path), "architecture"
    )

    assert [page.page_number for page in normalized.pages] == [1]
    assert normalized.pages[0].selection_reasons == [
        "all_pages_quality_pipeline"
    ]
    assert normalized.pages[0].image_path == "pages/page-0001.png"
    assert normalized.pages[0].vision_provenance.provider == "fake"
    assert normalized.pages[0].native_elements
    assert normalized.pages[0].reconciliation is not None
    assert normalized.pages[0].review is not None
    assert normalized.pages[0].review.status == "AUTO_PASS"
    assert "## Visual evidence" in normalized.pages[0].combined_markdown
    report = build_pdf_llm_report(normalized)
    assert "## Page 1" in report
    assert "![Rendered page 1](pages/page-0001.png)" in report
    assert "### Independent judge" in report
    assert "**AUTO_PASS**" in report


def test_pdf_material_discrepancy_preserves_reconciled_partial_output(
    tmp_path: Path,
) -> None:
    path = tmp_path / "review.pdf"
    document = canvas.Canvas(str(path))
    document.drawString(72, 750, "Native value 42")
    document.save()

    with pytest.raises(ExtractionFailure) as raised:
        hybrid_pipeline(material_discrepancy=True).extract(
            path, path.name, "official", digest(path), "general_document"
        )

    assert raised.value.code == "PDF_REVIEW_REQUIRED"
    assert raised.value.partial_document is not None
    page = raised.value.partial_document.pages[0]
    assert page.review is not None
    assert page.review.status == "REVIEW_REQUIRED"
    assert page.reconciliation is not None
    assert page.reconciliation.conflicts == ["A material wording mismatch"]


def test_pymupdf_renderer_creates_stable_selected_page_pngs(tmp_path: Path) -> None:
    path = tmp_path / "render.pdf"
    document = canvas.Canvas(str(path))
    document.drawString(72, 750, "Page one")
    document.showPage()
    document.drawString(72, 750, "Page two")
    document.save()

    renderer = PyMuPdfRenderer(dpi=144)
    first = renderer.render(path, 2, [2])
    second = renderer.render(path, 2, [2])

    assert [item.page_number for item in first] == [2]
    assert first[0].path.is_file()
    assert first[0].content_hash == second[0].content_hash


def test_blank_pdf_reports_ocr_required(tmp_path: Path) -> None:
    path = tmp_path / "scan.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=612, height=792)
    with path.open("wb") as handle:
        writer.write(handle)

    with pytest.raises(ExtractionFailure) as raised:
        extract(path)

    assert raised.value.code == "OCR_REQUIRED"
    assert raised.value.partial_document is not None


def test_encrypted_pdf_fails_clearly(tmp_path: Path) -> None:
    path = tmp_path / "encrypted.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=612, height=792)
    writer.encrypt("secret", algorithm="RC4-128")
    with path.open("wb") as handle:
        writer.write(handle)

    with pytest.raises(ExtractionFailure) as raised:
        extract(path)

    assert raised.value.code == "ENCRYPTED_DOCUMENT"


def test_unsupported_format_fails_clearly(tmp_path: Path) -> None:
    path = tmp_path / "image.png"
    path.write_bytes(b"not an accepted document")

    with pytest.raises(ExtractionFailure) as raised:
        extract(path)

    assert raised.value.code == "UNSUPPORTED_FORMAT"


@pytest.mark.parametrize("suffix", [".pdf", ".pptx", ".xlsx", ".docx"])
def test_corrupt_structured_document_fails_clearly(
    tmp_path: Path, suffix: str
) -> None:
    path = tmp_path / f"corrupt{suffix}"
    path.write_bytes(b"not a valid structured document")

    with pytest.raises(ExtractionFailure) as raised:
        extract(path)

    assert raised.value.code == "CORRUPT_DOCUMENT"


def test_empty_text_fails_clearly(tmp_path: Path) -> None:
    path = tmp_path / "empty.txt"
    path.write_text("   \n", encoding="utf-8")

    with pytest.raises(ExtractionFailure) as raised:
        extract(path)

    assert raised.value.code == "EMPTY_DOCUMENT"
