from __future__ import annotations

import hashlib
import os
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from app.extraction.docx_validation import (
    DocxPackageLimits,
    validate_docx_package,
)
from app.extraction.errors import ExtractionFailure
from app.extraction.pipeline import ExtractionPipeline
from docx import Document


def _digest(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _partial_native(path: Path, *, revision_policy: str = "final"):
    pipeline = ExtractionPipeline(docx_revision_policy=revision_policy)
    with pytest.raises(ExtractionFailure) as raised:
        pipeline.extract(path, path.name, "official", _digest(path))
    assert raised.value.code == "DOCX_RENDERER_REQUIRED"
    assert raised.value.partial_document is not None
    return raised.value.partial_document


def _rewrite_document_xml(path: Path, insertion: str) -> None:
    replacement = path.with_suffix(".rewritten.docx")
    with ZipFile(path) as source, ZipFile(
        replacement, "w", compression=ZIP_DEFLATED
    ) as target:
        for item in source.infolist():
            payload = source.read(item.filename)
            if item.filename == "word/document.xml":
                xml = payload.decode("utf-8")
                payload = xml.replace("</w:body>", insertion + "</w:body>").encode(
                    "utf-8"
                )
            target.writestr(item, payload)
    os.replace(replacement, path)


def test_docx_unwraps_content_controls_and_applies_final_revision_view(
    tmp_path: Path,
) -> None:
    path = tmp_path / "controlled.docx"
    document = Document()
    document.add_paragraph("Stable text")
    document.save(path)
    _rewrite_document_xml(
        path,
        """
        <w:sdt><w:sdtPr/><w:sdtContent><w:p><w:r><w:t>Controlled decision</w:t>
        </w:r></w:p></w:sdtContent></w:sdt>
        <w:p><w:ins w:id="1"><w:r><w:t>Final wording</w:t></w:r></w:ins>
        <w:del w:id="2"><w:r><w:delText>Original wording</w:delText></w:r>
        </w:del></w:p>
        """,
    )

    normalized = _partial_native(path, revision_policy="final")
    text = "\n".join(block.text for block in normalized.blocks)

    assert "Controlled decision" in text
    assert "Final wording" in text
    assert "Original wording" not in text
    assert normalized.docx_metadata.content_control_count == 1
    assert normalized.docx_metadata.inserted_revision_count == 1
    assert normalized.docx_metadata.deleted_revision_count == 1


def test_docx_original_revision_view_excludes_inserted_text(tmp_path: Path) -> None:
    path = tmp_path / "original.docx"
    document = Document()
    document.add_paragraph("Stable text")
    document.save(path)
    _rewrite_document_xml(
        path,
        """
        <w:p><w:ins w:id="1"><w:r><w:t>Final wording</w:t></w:r></w:ins>
        <w:del w:id="2"><w:r><w:delText>Original wording</w:delText></w:r>
        </w:del></w:p>
        """,
    )

    normalized = _partial_native(path, revision_policy="original")
    text = "\n".join(block.text for block in normalized.blocks)

    assert "Original wording" in text
    assert "Final wording" not in text


def test_docx_nested_tables_are_separate_auditable_blocks(tmp_path: Path) -> None:
    path = tmp_path / "nested.docx"
    document = Document()
    table = document.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "Outer value"
    nested = table.cell(0, 0).add_table(rows=1, cols=1)
    nested.cell(0, 0).text = "Nested value"
    document.save(path)

    normalized = _partial_native(path)
    tables = [block for block in normalized.blocks if block.kind == "table"]

    assert len(tables) == 2
    assert tables[0].attributes["parent_table"] is None
    assert tables[1].attributes["parent_table"] == 1
    assert "Nested value" in tables[1].text


def test_docx_package_limits_are_enforced_before_extraction(tmp_path: Path) -> None:
    path = tmp_path / "large.docx"
    document = Document()
    document.add_paragraph("First")
    document.add_paragraph("Second")
    document.save(path)

    with pytest.raises(ExtractionFailure) as raised:
        validate_docx_package(path, DocxPackageLimits(max_paragraphs=1))

    assert raised.value.code == "DOCX_LIMIT_EXCEEDED"


def test_docx_unsafe_zip_member_is_rejected(tmp_path: Path) -> None:
    path = tmp_path / "unsafe.docx"
    document = Document()
    document.add_paragraph("Safe native text")
    document.save(path)
    replacement = tmp_path / "unsafe-rewritten.docx"
    with ZipFile(path) as source, ZipFile(
        replacement, "w", compression=ZIP_DEFLATED
    ) as target:
        for item in source.infolist():
            target.writestr(item, source.read(item.filename))
        target.writestr("../escape.txt", b"unsafe")
    os.replace(replacement, path)

    with pytest.raises(ExtractionFailure) as raised:
        validate_docx_package(path, DocxPackageLimits())

    assert raised.value.code == "CORRUPT_DOCUMENT"
    assert "unsafe ZIP path" in str(raised.value)


def test_docx_active_content_is_rejected_before_rendering(tmp_path: Path) -> None:
    path = tmp_path / "active.docx"
    document = Document()
    document.add_paragraph("Safe native text")
    document.save(path)
    replacement = tmp_path / "active-rewritten.docx"
    with ZipFile(path) as source, ZipFile(
        replacement, "w", compression=ZIP_DEFLATED
    ) as target:
        for item in source.infolist():
            target.writestr(item, source.read(item.filename))
        target.writestr("word/vbaProject.bin", b"not-executed")
    os.replace(replacement, path)

    with pytest.raises(ExtractionFailure) as raised:
        validate_docx_package(path, DocxPackageLimits())

    assert raised.value.code == "UNSAFE_DOCUMENT"
    assert "active content" in str(raised.value)
